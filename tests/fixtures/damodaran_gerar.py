"""Gera, em memória, planilhas SINTÉTICAS que imitam o layout das planilhas
Damodaran (linhas de metadados + cabeçalho + algumas linhas de dados), com
números inventados. Não são cópias nem recortes dos arquivos reais do Damodaran
— apenas o formato (metadados no topo, cabeçalho, colunas) é reproduzido, para
que ``tests/test_damodaran.py`` possa validar o parser offline sem versionar
nenhum dado da fonte real no repositório (que é público).

Usadas apenas em teste; não fazem parte do pacote ``wacc_toolkit``.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook


def _para_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def planilha_beta_setor(data_atualizacao: datetime | None = datetime(2026, 1, 5)) -> bytes:
    """Imita betaGlobal.xls/betaemerg.xls/betas.xls: 9 linhas de metadados, depois
    cabeçalho "Industry Name, Number of firms, Beta , D/E Ratio, Effective Tax
    rate, Unlevered beta, Cash/Firm value, Unlevered beta corrected for cash",
    seguido de colunas extras variáveis por ano (que o parser deve ignorar) e
    3 setores fictícios."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Industry Averages"
    linhas_meta = [
        ["Date updated:", data_atualizacao],
        ["Created by:", "Fulano de Tal, ficticio@example.com"],
        ["What is this data?", "Beta sintético de teste"],
        ["Home Page:", "http://example.com"],
        ["Data website:", "http://example.com/data"],
        ["Companies in each industry:", "http://example.com/indname.xls"],
        ["Variable definitions:", "http://example.com/variable.htm"],
        ["Do you want to use marginal or effective tax rates?", None, None, None, None, "Marginal"],
        ["If marginal tax rate, enter the marginal tax rate to use", None, None, None, None, 0.21],
    ]
    for linha in linhas_meta:
        ws.append(linha)
    ws.append([
        "Industry Name", "Number of firms", "Beta ", "D/E Ratio", "Effective Tax rate",
        "Unlevered beta", "Cash/Firm value", "Unlevered beta corrected for cash",
        "HiLo Risk", "2024", "2025",
    ])
    setores = [
        ("Setor Alfa", 42, 1.10, 0.35, 0.21, 0.85, 0.05, 0.82, "baixo", 1.05, 1.08),
        ("Setor Beta", 15, 0.90, 0.10, 0.19, 0.83, 0.12, 0.75, "alto", 0.88, 0.91),
        ("Setor Gama", 7, 1.35, 0.60, 0.25, 0.95, 0.20, 0.79, "baixo", 1.30, 1.34),
    ]
    for linha in setores:
        ws.append(list(linha))
    return _para_bytes(wb)


def planilha_beta_setor_sem_data() -> bytes:
    """Mesma tabela, mas sem nenhuma linha "Date updated"/"Date of update" —
    usada para testar a falha explícita ao extrair a versão."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Industry Averages"
    ws.append(["Created by:", "Fulano de Tal"])
    ws.append(["What is this data?", "Beta sintético de teste"])
    ws.append([
        "Industry Name", "Number of firms", "Beta ", "D/E Ratio", "Effective Tax rate",
        "Unlevered beta", "Cash/Firm value", "Unlevered beta corrected for cash",
    ])
    ws.append(["Setor Alfa", 42, 1.10, 0.35, 0.21, 0.85, 0.05, 0.82])
    return _para_bytes(wb)


def planilha_totalbeta_setor(data_atualizacao: datetime | None = datetime(2026, 1, 5)) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Industry Averages"
    ws.append(["Date updated:", data_atualizacao])
    ws.append(["Created by:", "Fulano de Tal"])
    ws.append(["What is this data?", "Total beta sintético de teste"])
    ws.append(["Home Page:", "http://example.com"])
    ws.append(["Data website:", "http://example.com/data"])
    ws.append(["Companies in each industry:", "http://example.com/indname.xls"])
    ws.append(["Variable definitions:", "http://example.com/variable.htm"])
    ws.append([
        "Industry Name", "Number of firms", "Average Unlevered Beta", "Average Levered Beta",
        "Average correlation with the market", "Total Unlevered Beta", "Total Levered Beta",
    ])
    setores = [
        ("Setor Alfa", 42, 0.85, 1.10, 0.20, 4.25, 5.50),
        ("Setor Beta", 15, 0.83, 0.90, 0.15, 5.53, 6.00),
    ]
    for linha in setores:
        ws.append(list(linha))
    return _para_bytes(wb)


def planilha_dbtfund_setor(data_atualizacao: datetime | None = datetime(2026, 1, 5)) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Industry Averages"
    ws.append(["Date updated:", data_atualizacao])
    ws.append(["Created by:", "Fulano de Tal"])
    ws.append(["What is this data?", "Dívida sintética de teste"])
    ws.append(["Home Page:", "http://example.com"])
    ws.append(["Data website:", "http://example.com/data"])
    ws.append(["Companies in each industry:", "http://example.com/indname.xls"])
    ws.append(["Variable definitions:", "http://example.com/variable.htm"])
    ws.append([
        "Industry Name", "Number of firms", "Book Debt to Capital", "Market Debt to Capital (Unadjusted)",
        "Market D/E (unadjusted)", "Market Debt to Capital (adjusted for leases)",
        "Market D/E (adjusted for leases)", "Interest Coverage Ratio", "Debt to EBITDA",
        "Effective tax rate", "Institutional Holdings", "Std dev in Stock Prices", "EBITDA/EV",
        "Net PP&E/Total Assets", "Capital Spending/Total Assets",
    ])
    setores = [
        ("Setor Alfa", 42, 0.40, 0.35, 0.55, 0.36, 0.56, 8.5, 2.1, 0.21, 0.60, 0.30, 0.12, 0.45, 0.05),
        ("Setor Beta", 15, 0.20, 0.10, 0.11, 0.11, 0.12, -1.2, -9000.0, 0.19, 0.55, 0.20, 0.08, 0.30, 0.03),
    ]
    for linha in setores:
        ws.append(list(linha))
    return _para_bytes(wb)


def planilha_ctryprem(data_atualizacao: datetime | None = datetime(2026, 1, 1)) -> bytes:
    """Imita ctryprem.xlsx, aba "ERPs by country": cabeçalho da 2a coluna vem
    preenchido com o nome de um valor ("Regiao Exemplo") em vez de "Region" (o
    mesmo defeito observado na planilha real), tabela principal seguida de uma
    segunda tabela solta ("Frontier Markets...") que o parser deve descartar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ERPs by country"
    ws.append(["Country and Equity Risk Premiums"])
    ws.append(["Date of update:", data_atualizacao, "(nota qualquer)"])
    ws.append(["Enter the current risk premium for a mature equity market", None, None, None, 0.045])
    ws.append(["Enter the current risk premium for the US =", None, None, None, 0.046])
    ws.append(["ajustar pela volatilidade?", None, None, None, "Yes"])
    ws.append(["multiplicador", None, None, None, 1.4])
    ws.append([None])
    ws.append(["Country", "Regiao Exemplo", "Moody's rating", "Rating-based Default Spread",
               "Total Equity Risk Premium", "Country Risk Premium"])
    paises = [
        ("Paisalandia", "Regiao A", "Aa2", 0.004, 0.048, 0.006),
        ("Paisonia", "Regiao B", "Ba3", 0.030, 0.088, 0.046),
        ("Terra do Nunca", "Regiao A", "Caa1", 0.063, 0.139, 0.097),
    ]
    for linha in paises:
        ws.append(list(linha))
    ws.append(["Frontier Markets (no sovereign ratings)"])
    ws.append(["Country", "PRS Composite Risk Score", "ERP", "CRP", "Default Spread"])
    ws.append(["Pais Fronteira", 65, 0.11, 0.07, 0.05])
    return _para_bytes(wb)


def planilha_histretsp(data_atualizacao: datetime | None = datetime(2026, 1, 1)) -> bytes:
    """Imita histretSP.xls, aba "Returns by year": alguns anos de retornos
    fictícios seguidos das linhas de resumo ("Arithmetic Average...") que o
    parser deve descartar pelo filtro de ano de 4 dígitos."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Returns by year"
    ws.append(["Date updated:", data_atualizacao])
    ws.append(["Created by:", "Fulano de Tal"])
    ws.append(["What is this data?", "Retornos sintéticos de teste"])
    for _ in range(15):
        ws.append([None])
    ws.append(["Year", "S&P 500 (includes dividends)", "US Small cap (bottom decile)",
               "3-month T.Bill", "US T. Bond (10-year)", "Baa Corporate Bond"])
    anos = [
        (2023, 0.18, 0.10, 0.05, 0.03, 0.06),
        (2024, -0.05, -0.08, 0.05, -0.02, 0.01),
        (2025, 0.12, 0.09, 0.045, 0.04, 0.05),
    ]
    for linha in anos:
        ws.append(list(linha))
    ws.append([None])
    ws.append(["Arithmetic Average Historical Return"])
    ws.append(["1928-2025", 0.115, 0.14, 0.033, 0.05, 0.062])
    return _para_bytes(wb)
