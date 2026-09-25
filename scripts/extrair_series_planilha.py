#!/usr/bin/env python
"""Importação inicial: extrai CDS Brasil (10a, 5a) e Ibovespa de uma planilha WACC legada
e grava CSVs no formato "Dados Históricos" do Investing.com, prontos para o coletor
``wacc_toolkit.collectors.investing`` (basta copiá-los para ``<bases>/entrada/investing/``).

Não roda automaticamente — é executado manualmente pelo revisor, uma única vez, para
importar o histórico que hoje só existe dentro da planilha.

Uso::

    python scripts/extrair_series_planilha.py <planilha.xlsx> <pasta_saida>

A planilha é aberta com ``openpyxl`` em modo ``read_only=True, data_only=True`` — o
script NUNCA grava nem salva a planilha de origem, apenas lê.

Procedência dos dados (validado contra
``PPP_Escolas_Santa_Maria-WACC-v02.xlsx``, abas abaixo):

- Aba **"3.CDS 10Y"**: cabeçalho na linha 5 (colunas B–H: "Ano", "Data", "Último",
  "Abertura", "Máxima", "Mínima", "Var%"). Fonte citada na própria aba:
  br.investing.com/rates-bonds/brazil-cds-10-years-usd-historical-data.
- Aba **"3.CDS 5Y"**: mesmo layout (cabeçalho na linha 5, colunas B–H), fonte
  br.investing.com/rates-bonds/brazil-cds-5-years-usd-historical-data.
- Aba **"3.Ibov"**: cabeçalho na linha 6 (colunas B–J: "Ano", "Data CSV", "Data",
  "Último", "Abertura", "Máxima", "Mínima", "Var%", "LN RET"). Fontes citadas: B3 e
  br.investing.com/indices/bovespa-historical-data.

Em todas as abas, este script usa apenas as colunas "Data", "Último", "Abertura",
"Máxima" e "Mínima" (identificadas pelo texto do cabeçalho, não pela posição fixa —
tolera pequenas mudanças de layout). "Ano", "Var%", "LN RET" e as tabelas auxiliares
de médias (colunas mais à direita) são ignoradas.

ATENÇÃO — granularidade: na planilha inspecionada, as abas "3.CDS 10Y" e "3.CDS 5Y"
trazem uma linha por MÊS (ex.: 2026-01-01, 2025-12-01, ...), não por dia — aparentam
ser o histórico "mensal" baixado do Investing.com, diferente do CSV diário que o
coletor `investing` espera do usuário para as coletas correntes. Este script grava
os dados tal como estão na planilha (mensal); ao serem processados pelo coletor, isso
vai gerar avisos de "lacuna de dias" (``max_lagum_dias=15``) nesse trecho histórico —
é um aviso, não erro, e não bloqueia a gravação. A aba "3.Ibov" já é diária.

Os CSVs gerados usam o mesmo layout dos exports do Investing.com em português:
cabeçalho ``"Data","Último","Abertura","Máxima","Mínima","Vol.","Var%"``, datas
``dd.mm.aaaa``, números em formato brasileiro. Como a planilha não traz volume nem
sempre uma variação numérica confiável, essas duas colunas são gravadas como ``"-"``
(o coletor as ignora de qualquer forma). Nome do arquivo:
``<prefixo>_planilha_<aaaammdd-ini>_<aaaammdd-fim>.csv``, onde ``<prefixo>`` é
``cds10``, ``cds5`` ou ``ibovespa`` — escolhido para bater com as regras de
identificação por nome de arquivo do coletor `investing` (precisa conter "cds"+"10",
"cds"+"5", ou "ibovespa"/"bovespa", respectivamente).
"""

from __future__ import annotations

import datetime as dt
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

_ABAS = (
    ("3.CDS 10Y", "cds10"),
    ("3.CDS 5Y", "cds5"),
    ("3.Ibov", "ibovespa"),
)

_CABECALHOS_NECESSARIOS = {"data", "ultimo", "abertura", "maxima", "minima"}


def _sem_acento(texto: str) -> str:
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def _norm(valor: object) -> str:
    return _sem_acento(str(valor)).strip().lower() if valor is not None else ""


def localizar_cabecalho(ws, linhas_max: int = 15) -> tuple[int, dict[str, int]]:
    """Acha a linha de cabeçalho e a posição (índice 0-based na linha) de cada coluna
    necessária, procurando pelo texto (não pela posição fixa)."""
    for num, linha in enumerate(ws.iter_rows(min_row=1, max_row=linhas_max, values_only=True), start=1):
        posicoes: dict[str, int] = {}
        for idx, valor in enumerate(linha):
            texto = _norm(valor)
            if texto and texto not in posicoes:
                posicoes[texto] = idx
        if _CABECALHOS_NECESSARIOS.issubset(posicoes.keys()):
            return num, posicoes
    raise ValueError(
        f"{ws.title}: cabeçalho com Data/Último/Abertura/Máxima/Mínima não encontrado "
        f"nas primeiras {linhas_max} linhas"
    )


def extrair_serie(ws) -> pd.DataFrame:
    """Lê todas as linhas de dados de uma aba (após o cabeçalho) e devolve
    ``data,ultimo,abertura,maxima,minima``. Linhas sem data válida ou sem "Último"
    numérico são descartadas (rodapé, linhas em branco, "#DIV/0!" etc.)."""
    linha_cab, pos = localizar_cabecalho(ws)

    def _num(linha: tuple, chave: str) -> float | None:
        idx = pos[chave]
        v = linha[idx] if idx < len(linha) else None
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

    registros = []
    for linha in ws.iter_rows(min_row=linha_cab + 1, values_only=True):
        idx_data = pos["data"]
        data = linha[idx_data] if idx_data < len(linha) else None
        if not isinstance(data, (dt.datetime, dt.date)):
            continue
        ultimo = _num(linha, "ultimo")
        if ultimo is None:
            continue
        registros.append(
            {
                "data": data.date() if isinstance(data, dt.datetime) else data,
                "ultimo": ultimo,
                "abertura": _num(linha, "abertura"),
                "maxima": _num(linha, "maxima"),
                "minima": _num(linha, "minima"),
            }
        )

    if not registros:
        raise ValueError(f"{ws.title}: nenhuma linha de dados válida encontrada")

    df = pd.DataFrame(registros).drop_duplicates(subset=["data"], keep="last")
    return df.sort_values("data", kind="stable").reset_index(drop=True)


def _numero_br(v: float | None) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "-"
    s = f"{v:,.2f}"  # "1,234.56"
    return s.replace(",", "@").replace(".", ",").replace("@", ".")  # -> "1.234,56"


def gravar_csv_investing(df: pd.DataFrame, caminho: Path) -> None:
    """Grava no formato do CSV 'Dados Históricos' do Investing.com (mais recente primeiro)."""
    linhas = ['"Data","Último","Abertura","Máxima","Mínima","Vol.","Var%"']
    for row in df.sort_values("data", ascending=False, kind="stable").itertuples(index=False):
        data_str = row.data.strftime("%d.%m.%Y")
        linhas.append(
            f'"{data_str}","{_numero_br(row.ultimo)}","{_numero_br(row.abertura)}",'
            f'"{_numero_br(row.maxima)}","{_numero_br(row.minima)}","-","-"'
        )
    caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8")


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("uso: python scripts/extrair_series_planilha.py <planilha.xlsx> <pasta_saida>")
        return 1

    planilha = Path(argv[1])
    saida = Path(argv[2])
    saida.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(planilha, read_only=True, data_only=True)
    try:
        for aba, prefixo in _ABAS:
            if aba not in wb.sheetnames:
                print(f"aviso: aba {aba!r} não encontrada em {planilha.name}, pulando")
                continue
            ws = wb[aba]
            df = extrair_serie(ws)
            ini = df["data"].min().strftime("%Y%m%d")
            fim = df["data"].max().strftime("%Y%m%d")
            nome = f"{prefixo}_planilha_{ini}_{fim}.csv"
            caminho = saida / nome
            gravar_csv_investing(df, caminho)
            print(
                f"{aba}: {len(df)} linhas, {df['data'].min()} a {df['data'].max()}, "
                f"último valor = {df['ultimo'].iloc[-1]} -> {caminho}"
            )
    finally:
        wb.close()  # leitura apenas: nunca chamar wb.save()

    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
